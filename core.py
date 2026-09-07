# -*- coding: utf-8 -*-
"""
core.py — Logic dùng chung cho BOT Công tác phí (CLI + Streamlit).

Tất cả hàm ở đây làm việc trên BYTES / đối tượng file, không đụng tới đường dẫn
cố định, không print, không sys.exit. Kết quả trả về trong dataclass Result.
"""
from __future__ import annotations

import io
import re
import html
import email
import email.header
import zipfile
import tempfile
import unicodedata
import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# ================== ĐỊA CHỈ / ĐƠN VỊ NGƯỜI MUA (mặc định: CPM VIỆT NAM) ==================
DIA_CHI_CHUAN = "Tầng 5, Số 231 – 233 Lê Thánh Tôn, Phường Bến Thành, TP Hồ Chí Minh, Việt Nam"
TEN_NGUOI_MUA = "CÔNG TY CỔ PHẦN CPM VIỆT NAM"
MST_NGUOI_MUA = "0305908974"

MON = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

DNTT_INV_ROW0, DNTT_INV_ROW_MAX, TRACUU_ROW0 = 27, 33, 3


# ------------------------------ helpers thuần -----------------------------------
def collapse(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip()


def parse_money(s):
    if s is None:
        return None
    s = re.sub(r"[^\d]", "", str(s).strip())
    return int(s) if s else None


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if v in (None, ""):
        return None
    s = str(v).strip().replace("-", "/").replace(".", "/")
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_number(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[.,\s]", "", str(v)))
    except ValueError:
        return None


def parse_month_year(v):
    """'9/2026' | '2026-09' | datetime | '9' -> (month, year) hoặc (None, None)."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.month, v.year
    s = str(v or "").strip()
    mo = re.search(r"\b(\d{1,2})\s*[/\-]\s*(\d{4})\b", s)
    if mo:
        return int(mo.group(1)), int(mo.group(2))
    mo = re.search(r"\b(\d{4})\s*[/\-]\s*(\d{1,2})\b", s)
    if mo:
        return int(mo.group(2)), int(mo.group(1))
    mo = re.match(r"\s*(\d{1,2})\s*$", s)
    if mo:
        return int(mo.group(1)), dt.date.today().year
    return None, None


def chia_ky_cong_tac(tu, den, n):
    """Chia kỳ [tu..den] (cả 2 đầu) thành n đoạn ngày liền nhau, đoạn dài dồn về cuối,
    đoạn cuối kết thúc đúng 'den'.  21/07–20/08 (31 ngày) / 4 hóa đơn -> 7,8,8,8."""
    n = max(1, n)
    tong = (den - tu).days + 1
    base, rem = divmod(tong, n)
    doan, cur = [], tu
    for i in range(n):
        dai = base + (1 if i >= n - rem else 0)
        end = den if i == n - 1 else cur + dt.timedelta(days=max(1, dai) - 1)
        doan.append((cur, end))
        cur = end + dt.timedelta(days=1)
    return doan


def mileage_amount(km):
    """Công thức bậc thang: 1000 km đầu × 1.500đ, phần vượt × 1.900đ."""
    if km < 1000:
        return round(km * 1500)
    return round(1000 * 1500 + (km - 1000) * 1900)


def tidy_url(u):
    u = (u or "").strip().rstrip("/#")
    if not u:
        return ""
    if not u.startswith("http"):
        u = "https://" + u
    if "pvoil.vn" in u:
        return "https://hoadon.pvoil.vn/Invoice/search#"
    if "vnpt-invoice.com.vn" in u:
        return u + "/"
    return u


def norm_addr(s):
    s = unicodedata.normalize("NFC", s or "").lower().strip()
    s = s.replace("–", "-").replace("—", "-")
    s = re.sub(r"\s*-\s*", "-", s)
    s = re.sub(r"tp\.?\s*hồ chí minh|thành phố hồ chí minh|tp\s*hcm|tphcm|hcmc", "hcm", s)
    s = re.sub(r"\bsố\b", " ", s)
    s = s.replace(".", " ").replace(",", " ")
    return re.sub(r"\s+", " ", s).strip()


# ------------------------------ đọc hóa đơn -----------------------------------
def _pdf_text(pdf_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(pdf_bytes))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    except Exception:
        return ""


def extract_invoice(pdf_bytes: bytes, buyer=None) -> dict:
    buyer = buyer or {}
    ten_mua = buyer.get("ten", TEN_NGUOI_MUA)
    mst_mua = buyer.get("mst", MST_NGUOI_MUA)
    dc_chuan = buyer.get("dia_chi", DIA_CHI_CHUAN)

    t = collapse(_pdf_text(pdf_bytes))
    d = {"so_hd": None, "so_tien": None, "vat": 0, "web": "", "ma_tra_cuu": "",
         "mst_ban": "", "dia_chi_mua": "", "dia_chi_ok": False, "ngay": None,
         "loai": "X", "is_hotel": False, "so_dem": None}
    if not t:
        return d

    # ----- khách sạn / nhà nghỉ? -----
    if re.search(r"thuê\s*phòng|phòng\s*nghỉ|lưu\s*trú|khách\s*sạn|nhà\s*nghỉ|homestay|motel|hotel",
                 t, re.I):
        d["is_hotel"] = True
        d["loai"] = "KS"
        m = re.search(r"\bĐêm\s+(\d{1,2})\b", t) or re.search(r"phòng[^0-9]{0,30}?(\d{1,2})\s*đêm", t, re.I)
        if m:
            d["so_dem"] = int(m.group(1))

    m = (re.search(r"Số\s*\(No\)?\s*:\s*(\d{5,10})", t)
         or re.search(r"Ký hiệu[^0-9]{0,20}?Số:\s*(\d{5,10})", t)
         or re.search(r"\bSố:\s*(\d{5,10})\b", t))
    if m:
        d["so_hd"] = m.group(1)

    m = (re.search(r"[Tt]ra cứu[^:]*tại[^:]*:\s*(https?://[^\s,)]+|[a-z0-9.\-]+\.vn)", t)
         or re.search(r"[Ww]ebsite tra cứu[^:]*:\s*(https?://[^\s,)]+|[a-z0-9.\-]+\.vn)", t))
    if m:
        d["web"] = tidy_url(m.group(1))
    m = re.search(r"[Mm]ã tra cứu[^:]*:\s*([A-Za-z0-9_\-]+)", t)
    if m:
        d["ma_tra_cuu"] = m.group(1)

    if d["is_hotel"]:
        m = (re.search(r"Căn cước công dân\s*:\s*(\d{9,13})", t)
             or re.search(r"MST\s*/?\s*CCCD[^:]*:\s*(\d{9,13})", t))
        if m and m.group(1) != mst_mua:
            d["mst_ban"] = m.group(1)
    else:
        m = re.search(r"Đơn vị bán hàng[^:]*:\s*(.+?)\s*Mã số thuế[^:]*:\s*(\d{10,13})", t)
        if m:
            d["mst_ban"] = m.group(2)
        else:
            m = re.search(r"(SƠN HẢI|XĂNG DẦU[^0-9]{0,40})\s*(\d{10})", t)
            if m:
                d["mst_ban"] = m.group(2)

    kkknt = "KKKNT" in t or "KKKNT" in t.replace(" ", "")
    m = (re.search(r"Tổng cộng tiền thanh toán[^:]*:\s*([\d.]+)", t)
         or re.search(r"Tổng số tiền thanh toán[^:]*:\s*([\d.]+)", t)
         or re.search(r"Cộng tiền bán hàng[^:]*:\s*([\d.]+)", t)      # hóa đơn bán hàng (khách sạn)
         or re.search(r"Cộng tiền hàng[^:]*:\s*([\d.]+)", t))
    if m:
        d["so_tien"] = parse_money(m.group(1))
    if not kkknt:
        m = re.search(r"Tiền thuế GTGT[^:]*:\s*([\d.]+)", t)
        d["vat"] = (parse_money(m.group(1)) or 0) if m else 0

    m = re.search(r"Ngày\s*(\d{1,2})\s*tháng\s*(\d{1,2})\s*năm\s*(\d{4})", t)
    if m:
        try:
            d["ngay"] = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            pass

    if not d["is_hotel"]:
        m = re.search(r"(Xăng|Dầu)\s+[^\n]{0,40}?(Lít|Kg|lít|kg)\b", t, re.I)
        if m:
            d["loai"] = "D" if m.group(1)[0].lower() == "d" else "X"

    # địa chỉ người mua: khối "Tầng N ... Việt Nam" (bên bán là cây xăng, không mở đầu "Tầng")
    m = re.search(r"(Tầng\s*\d+[^\n]{5,120}?(?:Việt\s*Nam|Vietnam)\.?)", t, re.I)
    if m:
        d["dia_chi_mua"] = collapse(m.group(1))
    d["dia_chi_ok"] = (
        ten_mua.lower() in t.lower()
        and mst_mua in t
        and bool(d["dia_chi_mua"])
        and norm_addr(d["dia_chi_mua"]) == norm_addr(dc_chuan)
    )
    return d


def _eml_text(eml_bytes: bytes):
    try:
        m = email.message_from_bytes(eml_bytes)
    except Exception:
        return "", ""
    subj = str(email.header.make_header(email.header.decode_header(m.get("Subject", ""))))
    body = ""
    for p in m.walk():
        if p.get_content_type() in ("text/plain", "text/html"):
            try:
                body += p.get_payload(decode=True).decode(
                    p.get_content_charset() or "utf-8", errors="replace")
            except Exception:
                pass
    return subj, collapse(html.unescape(re.sub(r"<[^>]+>", " ", body)))


def eml_invoice_no(eml_bytes: bytes):
    subj, text = _eml_text(eml_bytes)
    for pat in (r"[Ss]ố hóa đơn[:\s]*(\d{4,10})",
                r"[Hh]óa đơn điện tử số[:\s]*(\d{4,10})",
                r"gửi hóa đơn[^0-9]{0,20}số\s*(\d{4,10})",
                r"\bSố[:\s]*(\d{5,10})\b"):
        mm = re.search(pat, subj + " " + text)
        if mm:
            return mm.group(1)
    return None


def eml_info(eml_bytes: bytes) -> dict:
    """Rút số HĐ + mã tra cứu + website từ EML (dùng khi PDF là ảnh không có chữ)."""
    subj, text = _eml_text(eml_bytes)
    both = subj + " " + text
    d = {"so_hd": eml_invoice_no(eml_bytes), "ma_tra_cuu": "", "web": "", "is_hotel": None}
    m = re.search(r"[Mm]ã\s*(?:số|tra cứu)[^:]{0,12}[:\s]+([A-Za-z0-9_\-*]{6,})", both)
    if m:
        d["ma_tra_cuu"] = m.group(1)
    m = re.search(r"(https?://[a-z0-9.\-]+(?:tra-cuu|invoice|hoadon)[a-z0-9./\-]*)", both, re.I)
    if m:
        d["web"] = tidy_url(m.group(1))
    elif re.search(r"petrolimex", both, re.I):
        d["web"] = "https://hoadon.petrolimex.com.vn/"
    elif re.search(r"meinvoice|misa", both, re.I):
        d["web"] = "https://www.meinvoice.vn/tra-cuu"
    if re.search(r"petrolimex|xăng|dầu", both, re.I):
        d["is_hotel"] = False
    elif re.search(r"phòng nghỉ|khách sạn|nhà nghỉ", both, re.I):
        d["is_hotel"] = True
    return d


# ------------------------------ KM từ Numbers -----------------------------------
def tinh_km(numbers_bytes: bytes, tu, den):
    """(tong_km, {ngày: km}, [sheet thiếu]).  Đọc bảng 'Route Plan' cột Date + Km."""
    from numbers_parser import Document
    with tempfile.NamedTemporaryFile(suffix=".numbers", delete=False) as tmp:
        tmp.write(numbers_bytes)
        tmp_path = tmp.name
    try:
        doc = Document(tmp_path)
        routes = {}
        for sh in doc.sheets:
            for tb in sh.tables:
                if tb.name.strip() == "Route Plan":
                    routes[sh.name] = tb

        thang = set()
        d = tu
        while d <= den:
            thang.add((d.month, d.year))
            d += dt.timedelta(days=1)

        tong, chi_tiet, thieu = 0.0, {}, []
        for (mm, yy) in sorted(thang):
            ten = f"Lịch đi shop {mm}/{yy % 100}"
            tb = routes.get(ten) or next((routes[k] for k in routes if k.strip() == ten), None)
            if tb is None:
                thieu.append(ten)
                continue
            for r, cells in enumerate(tb.rows(values_only=True)):
                if r == 0 or len(cells) < 7:
                    continue
                ngay, km = cells[2], cells[6]
                if isinstance(ngay, dt.datetime):
                    ngay = ngay.date()
                if isinstance(ngay, dt.date) and isinstance(km, (int, float)) and tu <= ngay <= den:
                    tong += km
                    chi_tiet[ngay] = chi_tiet.get(ngay, 0.0) + km
        return tong, chi_tiet, thieu
    finally:
        try:
            Path(tmp_path).unlink()
        except OSError:
            pass


# ------------------------------ dataclasses -----------------------------------
@dataclass
class Config:
    month: int
    year: int
    tu_ngay: dt.date
    den_ngay: dt.date
    dia_diem: str = ""
    muc_dich: str = ""
    so_dem: int = 0
    so_ngay_pd: int = 0
    tam_ung: int = 0
    km_override: float | None = None
    strict_addr: bool = False
    buyer: dict = field(default_factory=dict)   # {ten, mst, dia_chi} — để trống = CPM

    @property
    def thang_label(self):
        return f"{self.month}/{self.year}"


@dataclass
class Result:
    ok: bool = False
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    log: list = field(default_factory=list)
    km: float = 0.0
    km_source: str = ""
    km_detail: dict = field(default_factory=dict)
    tien_di_lai: int = 0
    tien_ks: int = 0
    tien_pd: int = 0
    tong_cp: int = 0
    thu_chi_them: int = 0
    invoices: list = field(default_factory=list)      # dict/hàng đưa vào Excel
    tong_hd: int = 0
    khop: bool = False
    thieu: int = 0
    xlsx_bytes: bytes | None = None
    out_filename: str = ""
    renamed: list = field(default_factory=list)       # (tên cũ, tên mới)


# ------------------------------ điền Excel -----------------------------------
def fill_workbook(template_bytes: bytes, cfg: Config, invoices, km) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))

    dn = wb["DNTT"]
    dn["G3"] = f"Thời gian/ Time: {cfg.thang_label}"
    c10 = str(dn["C10"].value or "Công tác phí tháng")
    c10 = re.sub(r"th[áa]ng\s*\S*\s*$", f"tháng {cfg.thang_label}", c10.strip())
    if "tháng" not in c10:
        c10 += f" tháng {cfg.thang_label}"
    dn["C10"] = c10
    for r in range(DNTT_INV_ROW0, DNTT_INV_ROW_MAX + 1):
        for col in ("B", "F", "G", "H", "I", "J"):
            dn[f"{col}{r}"].value = None
            dn[f"{col}{r}"].hyperlink = None
    for i, inv in enumerate(invoices):
        r = DNTT_INV_ROW0 + i
        if r > DNTT_INV_ROW_MAX:
            raise ValueError("Quá 7 hóa đơn — mẫu DNTT không đủ dòng.")
        dn[f"A{r}"] = i + 1
        dn[f"B{r}"] = inv["mota"]
        dn[f"F{r}"] = cfg.thang_label
        dn[f"G{r}"] = inv["chua_vat"]
        dn[f"H{r}"] = inv["vat"]
        dn[f"I{r}"] = inv["thanh_toan"]
        dn[f"J{r}"] = inv["so_hd"]

    bk = wb[next(s for s in wb.sheetnames if s.strip() == "Bảng kê chi tiết")]
    bk["H3"] = f"Thời gian/ Time: {cfg.month:02d}/{cfg.year}"
    if cfg.muc_dich:
        bk["C4"] = cfg.muc_dich
    bk["C7"] = cfg.dia_diem
    bk["D7"] = (f"{cfg.tu_ngay.month}/{cfg.tu_ngay.day}/{cfg.tu_ngay.year} - "
                f"{cfg.den_ngay.month}/{cfg.den_ngay.day}/{cfg.den_ngay.year}")
    bk["E7"] = cfg.so_dem
    bk["G7"] = cfg.so_ngay_pd
    bk["D16"] = round(km)
    bk["H23"] = cfg.tam_ung if cfg.tam_ung else None

    tc = wb[next(s for s in wb.sheetnames if s.strip() == "Tra cứu HĐ")]
    for r in range(TRACUU_ROW0, TRACUU_ROW0 + 30):
        for col in ("B", "C", "D", "E"):
            tc[f"{col}{r}"].value = None
            tc[f"{col}{r}"].hyperlink = None
    for i, inv in enumerate(invoices):
        r = TRACUU_ROW0 + i
        tc[f"B{r}"] = inv["loai_hd"]
        tc[f"C{r}"] = inv["web"]
        tc[f"D{r}"] = inv["ma_tra_cuu"]
        tc[f"E{r}"] = inv["so_hd"]

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================== ORCHESTRATOR ==================================
def generate(template_bytes: bytes,
             invoice_files: list,          # [(filename, bytes), ...]  .pdf / .eml
             cfg: Config,
             manual_invoices: list | None = None,   # [{so_hd, mota, override}]
             numbers_bytes: bytes | None = None) -> Result:
    """Chạy toàn bộ quy trình. Không raise — mọi lỗi vào Result.errors."""
    R = Result()

    def err(m): R.errors.append(m)
    def warn(m): R.warnings.append(m)
    def log(m): R.log.append(m)

    # --- validate cơ bản ---
    if not template_bytes:
        err("Chưa có file Excel mẫu (tháng trước).")
    if cfg.den_ngay < cfg.tu_ngay:
        err(f"'Đến ngày' ({cfg.den_ngay:%d/%m/%Y}) sớm hơn 'Từ ngày' ({cfg.tu_ngay:%d/%m/%Y}).")
    pdfs = [(n, b) for n, b in invoice_files if n.lower().endswith(".pdf")]
    emls = [(n, b) for n, b in invoice_files if n.lower().endswith(".eml")]
    if not pdfs:
        err("Chưa upload file PDF hóa đơn nào.")
    if R.errors:
        return R

    # --- 1) đọc từng PDF ---
    log(f"Đọc {len(pdfs)} PDF + {len(emls)} EML.")
    data = {}            # so_hd -> info
    for name, b in pdfs:
        inv = extract_invoice(b, cfg.buyer)
        so = inv["so_hd"]
        if not so:
            warn(f"Không đọc được số hóa đơn trong: {name}")
            continue
        pre = "HĐD" if inv["loai"] == "D" else "HĐX"
        mark = "" if inv["dia_chi_ok"] else " [SAI ĐỊA CHỈ]"
        inv["_pre"] = pre
        new_name = f"{pre}_{so}{mark}.pdf"
        R.renamed.append((name, new_name))
        data[so] = inv
        if not inv["dia_chi_ok"]:
            warn(f"HĐ {so}: SAI/THIẾU ĐỊA CHỈ người mua — "
                 f"thấy “{inv['dia_chi_mua'] or 'không tìm thấy'}”.")

    if not data:
        err("Không PDF nào đọc được số hóa đơn.")
        return R

    for name, b in emls:
        so = (eml_invoice_no(b) or "").lstrip("0")
        match = next((k for k in data if k.lstrip("0") == so and so), None)
        if match:
            pre = data[match]["_pre"]
            mark = "" if data[match]["dia_chi_ok"] else " [SAI ĐỊA CHỈ]"
            R.renamed.append((name, f"{pre}_{match}{mark}.eml"))
        else:
            warn(f"EML không khớp hóa đơn nào: {name}")

    # --- 2) danh sách hóa đơn ---
    rows = []
    if manual_invoices:
        for m in manual_invoices:
            so = re.sub(r"\D", "", str(m.get("so_hd") or "")) or None
            mota = str(m.get("mota") or "").strip()
            if not so and not mota:
                continue
            rows.append({"so_hd": so, "mota": mota, "override": parse_money(m.get("override"))})
        # khớp với PDF
        miss = []
        for row in rows:
            pd_ = data.get(row["so_hd"]) or next(
                (v for k, v in data.items() if k.lstrip("0") == (row["so_hd"] or "").lstrip("0")), None)
            if pd_ is None:
                miss.append(row["so_hd"] or row["mota"])
            else:
                row["_pdf"] = pd_
        if miss:
            err("Hóa đơn trong danh sách KHÔNG thấy PDF: " + ", ".join(map(str, miss))
                + f"\nPDF đang có: {', '.join(sorted(data))}")
            return R
    else:
        # tự lập từ TẤT CẢ pdf: sắp theo ngày; trùng ngày -> tiền nhỏ trước
        items = list(data.values())
        for v in items:
            if v["ngay"] and not (cfg.tu_ngay <= v["ngay"] <= cfg.den_ngay):
                warn(f"HĐ {v['so_hd']} ghi ngày {v['ngay']:%d/%m/%Y} — ngoài kỳ "
                     f"{cfg.tu_ngay:%d/%m}–{cfg.den_ngay:%d/%m} (vẫn dùng).")
        order = sorted(items, key=lambda v: (v["ngay"] or dt.date(2100, 1, 1),
                                             v["so_tien"] or 0, v["so_hd"]))
        rows = [{"so_hd": v["so_hd"], "mota": "", "override": None, "_pdf": v} for v in order]
        log(f"Tự lập {len(rows)} hóa đơn (giữ hết), sắp theo ngày.")

    if not rows:
        err("Không có hóa đơn nào.")
        return R
    if len(rows) > 7:
        err(f"{len(rows)} hóa đơn — mẫu DNTT chỉ 7 dòng. Bớt bớt hóa đơn.")
        return R

    if cfg.strict_addr:
        bad = [r["so_hd"] for r in rows if not r["_pdf"]["dia_chi_ok"]]
        if bad:
            err("Chế độ 'Chính xác tuyệt đối': các HĐ sai địa chỉ — " + ", ".join(bad))
            return R

    # --- 2b) nội dung DNTT: chia kỳ thành N đoạn ngày ---
    doan = chia_ky_cong_tac(cfg.tu_ngay, cfg.den_ngay, len(rows))
    for row, (s, e) in zip(rows, doan):
        loai_hd = "HĐD" if row["_pdf"]["loai"] == "D" else "HĐX"
        loai_txt = "dầu" if loai_hd == "HĐD" else "xăng"
        if not row["mota"]:
            row["mota"] = f"CP {loai_txt} xe: {s:%d/%m/%Y} - {e:%d/%m/%Y}"
        row["loai_hd"] = loai_hd
        row["web"] = row["_pdf"]["web"]
        row["ma_tra_cuu"] = row["_pdf"]["ma_tra_cuu"]
        row["vat"] = int(round(row["_pdf"]["vat"] or 0))
        real = row["override"] if row["override"] is not None else row["_pdf"]["so_tien"]
        if real is None:
            err(f"Không đọc được số tiền HĐ {row['so_hd']}. Nhập tay 'Số tiền (ghi đè)'.")
            return R
        row["thanh_toan"] = int(round(real))

    # --- 3) KM + tiền ---
    if cfg.km_override:
        R.km = float(cfg.km_override)
        R.km_source = "nhập tay"
    elif numbers_bytes:
        try:
            tong, ct, missing = tinh_km(numbers_bytes, cfg.tu_ngay, cfg.den_ngay)
        except Exception as ex:
            err(f"Không đọc được file Numbers: {ex}")
            return R
        if missing:
            err("File Numbers thiếu sheet Route Plan cho: " + ", ".join(missing)
                + "\n-> Chọn 'Gõ tay số KM' và nhập tổng KM.")
            return R
        R.km, R.km_detail, R.km_source = tong, ct, "file Numbers"
    else:
        err("Chưa có số KM: upload file Numbers, hoặc chọn 'Gõ tay' và nhập tổng KM.")
        return R

    R.tien_di_lai = mileage_amount(R.km)
    R.tien_ks = cfg.so_dem * 500_000
    R.tien_pd = cfg.so_ngay_pd * 200_000
    R.tong_cp = R.tien_di_lai + R.tien_ks + R.tien_pd
    R.thu_chi_them = R.tong_cp - cfg.tam_ung

    # --- 4) hóa đơn cuối = số bù (chỉnh XUỐNG, không bịa tăng) ---
    for r in rows:
        r["chua_vat"] = r["thanh_toan"] - r["vat"]
    last = rows[-1]
    goc_last = last["thanh_toan"]
    tong_khac = sum(r["thanh_toan"] for r in rows[:-1])
    can_co = int(round(R.tien_di_lai - tong_khac))

    if can_co < 0:
        err(f"Các hóa đơn TRƯỚC hóa đơn cuối ({tong_khac:,.0f}) đã VƯỢT tiền đi lại theo KM "
            f"({R.tien_di_lai:,.0f}). Bớt hóa đơn hoặc kiểm tra lại KM / kỳ ngày.")
        return R
    if can_co < goc_last:
        last["thanh_toan"] = can_co
        last["chua_vat"] = can_co - last["vat"]
        log(f"Hóa đơn cuối HĐ {last['so_hd']}: {goc_last:,.0f} → {can_co:,.0f} "
            f"(chỉnh XUỐNG cho khớp KM).")
    elif can_co == goc_last:
        log("Tổng hóa đơn khớp sẵn tiền đi lại theo KM — không chỉnh gì.")
    else:
        R.thieu = can_co - goc_last
        warn(f"THIẾU HÓA ĐƠN: còn thiếu {R.thieu:,.0f} đồng so với tiền đi lại theo KM. "
             f"Bổ sung thêm hóa đơn rồi tạo lại. (BOT KHÔNG bịa tăng hóa đơn cuối.)")

    for j, r in enumerate(rows):
        r["is_last"] = (j == len(rows) - 1)
    R.invoices = rows
    R.tong_hd = sum(r["thanh_toan"] for r in rows)
    R.khop = abs(R.tong_hd - R.tien_di_lai) < 1

    # --- 5) xuất Excel ---
    try:
        R.xlsx_bytes = fill_workbook(template_bytes, cfg, rows, R.km)
    except Exception as ex:
        err(f"Lỗi khi điền Excel: {ex}")
        return R
    R.out_filename = f"CS Le Cong Bao Long - {MON[cfg.month]} {cfg.year}.xlsx"
    R.ok = True
    return R


def build_renamed_zip(invoice_files: list, renamed: list) -> bytes:
    """Đóng gói lại các file PDF/EML với tên mới (HĐX_<số>...) thành 1 zip."""
    by_name = {n: b for n, b in invoice_files}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for old, new in renamed:
            if old in by_name:
                z.writestr(new, by_name[old])
    return buf.getvalue()


def build_full_zip(result: "Result", invoice_files: list) -> bytes:
    """1 file zip chứa TẤT CẢ: file Excel công tác phí + các hóa đơn đã đổi tên."""
    by_name = {n: b for n, b in invoice_files}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        if result.xlsx_bytes:
            z.writestr(result.out_filename, result.xlsx_bytes)
        for old, new in result.renamed:
            if old in by_name:
                z.writestr(f"Hóa đơn/{new}", by_name[old])
    return buf.getvalue()


# ============================================================================
# ===================  CHẾ ĐỘ NHIỀU CHUYẾN / CÓ LƯU TRÚ  ======================
# ============================================================================

def _strip_diacritics(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s or "")
                   if unicodedata.category(c) != "Mn").replace("đ", "d").replace("Đ", "D")


@dataclass
class Trip:
    dia_diem: str
    tu: dt.date
    den: dt.date
    so_dem: int = 0
    km: float = 0.0

    @property
    def range_str(self):
        return f"{self.tu:%d/%m/%Y} - {self.den:%d/%m/%Y}"


@dataclass
class Layout:
    # DNTT
    dntt_desc_row0: int          # dòng diễn giải đầu (thường 10)
    dntt_n_desc: int             # số dòng diễn giải (1 hoặc 2)
    dntt_total_row: int          # dòng "Tổng cộng/ Total:"
    dntt_inv_row0: int           # dòng hóa đơn đầu trong "Bảng kê chứng từ"
    dntt_inv_slots: int          # số dòng hóa đơn tối đa
    dntt_sum_row: int
    # Bảng kê chi tiết
    bk_sheet: str
    bk_trip_row0: int
    bk_trip_slots: int
    bk_claim_row0: int
    bk_claim_slots: int
    bk_claim_total_row: int
    bk_h3: str = "H3"
    hotel_rate: int = 500_000
    pd_rate: int = 200_000
    bk_tong_cp_cell: str = "H25"
    bk_tam_ung_cell: str = "H23"
    bk_muc_dich_cell: str = "C4"
    # Tra cứu HĐ
    tc_sheet: str = "Tra cứu HĐ"
    tc_row0: int = 3
    requestor: str = ""


def _find_row(ws, col, needle, r1=1, r2=60):
    nl = needle.lower()
    for r in range(r1, r2):
        v = ws[f"{col}{r}"].value
        if v is not None and nl in str(v).lower():
            return r
    return None


def detect_layout(wb) -> Layout:
    dn = wb["DNTT"]
    bk_name = next(s for s in wb.sheetnames if s.strip() == "Bảng kê chi tiết")
    bk = wb[bk_name]
    tc_name = next((s for s in wb.sheetnames if s.strip() == "Tra cứu HĐ"), "Tra cứu HĐ")

    # --- DNTT: khối trên ---
    r_stt = _find_row(dn, "A", "STT", 1, 15) or 9
    desc0 = r_stt + 1
    total_row = _find_row(dn, "A", "Tổng cộng", desc0, desc0 + 8) or (desc0 + 1)
    n_desc = max(1, total_row - desc0)

    # --- DNTT: bảng kê chứng từ ---
    r_bk = _find_row(dn, "A", "BẢNG KÊ CHỨNG TỪ", 15, 40)
    hdr = (_find_row(dn, "A", "STT", r_bk + 1, r_bk + 5) if r_bk else None) or (r_bk + 2 if r_bk else 26)
    inv0 = hdr + 2                       # có 1 dòng phụ đề "(chưa VAT)" ở giữa
    sum_row = None
    for r in range(inv0, inv0 + 40):
        g = dn[f"G{r}"].value
        if isinstance(g, str) and g.replace(" ", "").upper().startswith("=SUM(G"):
            sum_row = r
            break
    if sum_row is None:
        sum_row = inv0 + 7
    slots = sum_row - inv0

    # --- Bảng kê chi tiết ---
    r_dd = _find_row(bk, "B", "Địa điểm công tác", 1, 12) or 5
    trip0 = r_dd + 2
    r_tong = _find_row(bk, "D", "Tổng:", trip0, trip0 + 15) or (trip0 + 6)
    trip_slots = max(1, r_tong - trip0)

    r_claim_hdr = _find_row(bk, "B", "(No.)", r_tong, r_tong + 6) or (r_tong + 3)
    claim0 = r_claim_hdr + 1
    r_claim_total = _find_row(bk, "B", "Tổng cộng (Total)", claim0, claim0 + 15) or (claim0 + 6)
    claim_slots = max(1, r_claim_total - claim0)

    r_tongcp = _find_row(bk, "E", "Tổng chi phí", r_claim_total, r_claim_total + 5) or (r_claim_total + 1)
    r_tamung = _find_row(bk, "E", "Tạm ứng", r_tongcp, r_tongcp + 4) or (r_tongcp + 1)
    r_mucdich = _find_row(bk, "B", "Mục đích công tác", 1, 8) or 4

    # định mức
    def _rate(label, default):
        rr = _find_row(bk, "H", label, 4, 10)
        if rr:
            for col in ("I", "J", "K"):
                v = bk[f"{col}{rr}"].value
                if isinstance(v, (int, float)) and v > 1000:
                    return int(v)
        return default
    hotel_rate = _rate("Khách sạn", 500_000)
    pd_rate = _rate("Per-diem", 200_000)

    h3 = "H3" if isinstance(bk["H3"].value, str) and "Thời gian" in str(bk["H3"].value) else \
         (f"H{_find_row(bk, 'H', 'Thời gian/ Time', 1, 6) or 3}")

    return Layout(
        dntt_desc_row0=desc0, dntt_n_desc=n_desc, dntt_total_row=total_row,
        dntt_inv_row0=inv0, dntt_inv_slots=slots, dntt_sum_row=sum_row,
        bk_sheet=bk_name, bk_trip_row0=trip0, bk_trip_slots=trip_slots,
        bk_claim_row0=claim0, bk_claim_slots=claim_slots, bk_claim_total_row=r_claim_total,
        bk_h3=h3, hotel_rate=hotel_rate, pd_rate=pd_rate,
        bk_tong_cp_cell=f"H{r_tongcp}", bk_tam_ung_cell=f"H{r_tamung}",
        bk_muc_dich_cell=f"C{r_mucdich}",
        tc_sheet=tc_name, tc_row0=(_find_row(wb[tc_name], "B", "Loại HĐ", 1, 6) or 2) + 1,
        requestor=str(dn["D5"].value or "").strip(),
    )


def _match_hotel_trip(ngay, trips):
    """Khớp hóa đơn KS với chuyến CÓ NGỦ LẠI: chọn chuyến ngắn nhất chứa ngày;
    nếu không có, chọn chuyến ngủ lại gần ngày nhất."""
    overnight = [(i, t) for i, t in enumerate(trips) if t.so_dem > 0] or list(enumerate(trips))
    if ngay:
        chua = [(i, t) for i, t in overnight
                if t.tu - dt.timedelta(days=1) <= ngay <= t.den + dt.timedelta(days=1)]
        if chua:
            return min(chua, key=lambda it: (it[1].den - it[1].tu).days)[0]
        return min(overnight,
                   key=lambda it: min(abs((ngay - it[1].tu).days), abs((ngay - it[1].den).days)))[0]
    return overnight[0][0]


def generate_multi(template_bytes: bytes,
                   invoice_files: list,
                   cfg: Config,
                   trips: list,               # list[Trip]  (đã parse)
                   gas_rows: list,             # [{so_hd, so_tien, chuyen}]  user nhập
                   khu_vuc: str = "") -> Result:
    """Chế độ nhiều chuyến / có lưu trú."""
    R = Result()
    err = R.errors.append; warn = R.warnings.append; log = R.log.append

    if not template_bytes:
        err("Chưa upload file Excel mẫu.")
    if not trips:
        err("Chưa nhập chuyến công tác nào.")
    for i, tr in enumerate(trips, 1):
        if not tr.dia_diem or not tr.tu or not tr.den:
            err(f"Chuyến {i}: thiếu địa điểm / ngày.")
        elif tr.den < tr.tu:
            err(f"Chuyến {i}: 'đến ngày' sớm hơn 'từ ngày'.")
    if R.errors:
        return R

    try:
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        L = detect_layout(wb)
    except Exception as ex:
        err(f"Không đọc được file Excel mẫu: {ex}")
        return R

    # --- đọc hóa đơn ---
    pdfs = [(n, b) for n, b in invoice_files if n.lower().endswith(".pdf")]
    emls = [(n, b) for n, b in invoice_files if n.lower().endswith(".eml")]
    pdf_by_no = {}       # so_hd(lstrip0) -> info
    for name, b in pdfs:
        inv = extract_invoice(b, cfg.buyer)
        so = inv["so_hd"]
        if so:
            pdf_by_no[so.lstrip("0") or so] = inv
            pre = "HĐKS" if inv["is_hotel"] else ("HĐD" if inv["loai"] == "D" else "HĐX")
            mark = "" if inv["dia_chi_ok"] else " [SAI ĐỊA CHỈ]"
            R.renamed.append((name, f"{pre}_{so}{mark}.pdf"))
            if not inv["dia_chi_ok"]:
                warn(f"HĐ {so}: địa chỉ người mua chưa khớp — thấy “{inv['dia_chi_mua'] or 'không tìm thấy'}”.")
    eml_by_no = {}
    for name, b in emls:
        info = eml_info(b)
        so = (info.get("so_hd") or "").lstrip("0")
        if so:
            eml_by_no[so] = info

    # --- KHÁCH SẠN: từ PDF, tự khớp chuyến ---
    hotel_lines = []
    for key, inv in pdf_by_no.items():
        if not inv["is_hotel"]:
            continue
        ti = _match_hotel_trip(inv["ngay"], trips)
        tr = trips[ti]
        hotel_lines.append({
            "so_hd": inv["so_hd"], "loai_hd": "HĐKS",
            "mota": f"Chi phí khách sạn {tr.range_str}",
            "chua_vat": int(round((inv["so_tien"] or 0) - (inv["vat"] or 0))),
            "vat": int(round(inv["vat"] or 0)),
            "thanh_toan": int(round(inv["so_tien"] or 0)),
            "web": inv["web"] or "https://www.meinvoice.vn/tra-cuu",
            "ma_tra_cuu": inv["ma_tra_cuu"], "mst_ban": inv["mst_ban"],
            "trip_idx": ti, "is_last": False,
        })
    if not hotel_lines:
        warn("Không thấy hóa đơn khách sạn nào trong file upload.")
    hotel_lines.sort(key=lambda h: h["trip_idx"])

    # --- XĂNG: từ bảng user nhập ---
    gas_lines = []
    for g in gas_rows:
        so = re.sub(r"\D", "", str(g.get("so_hd") or ""))
        tien = parse_money(g.get("so_tien"))
        if not so or tien is None:
            continue
        ch = g.get("chuyen")
        try:
            ti = int(ch) - 1
        except (TypeError, ValueError):
            ti = None
        if ti is None or not (0 <= ti < len(trips)):
            ti = len(gas_lines) % len(trips)     # rải đều nếu không ghi chuyến
        key = so.lstrip("0") or so
        pdf = pdf_by_no.get(key)
        eml = eml_by_no.get(key)
        web = (pdf and pdf["web"]) or (eml and eml.get("web")) or ""
        ma = (pdf and pdf["ma_tra_cuu"]) or (eml and eml.get("ma_tra_cuu")) or ""
        mst = (pdf and pdf["mst_ban"]) or ""
        loai_hd = "HĐD" if (pdf and pdf["loai"] == "D") else "HĐX"
        vat = int(round(pdf["vat"])) if (pdf and pdf["vat"]) else 0
        tr = trips[ti]
        gas_lines.append({
            "so_hd": so, "loai_hd": loai_hd,
            "mota": f"Chi phí {tr.dia_diem}: {tr.range_str}",
            "chua_vat": tien - vat, "vat": vat, "thanh_toan": tien,
            "web": web, "ma_tra_cuu": ma, "mst_ban": mst,
            "trip_idx": ti, "is_last": False,
        })
    if not gas_lines:
        err("Chưa nhập hóa đơn xăng nào (bảng 'Số tiền hóa đơn xăng').")
        return R
    gas_lines.sort(key=lambda x: (x["trip_idx"], x["so_hd"]))

    # --- tính tiền ---
    total_km = sum(t.km for t in trips)
    if total_km <= 0:
        err("Tổng KM = 0. Nhập KM cho các chuyến.")
        return R
    R.km = total_km
    R.km_source = "nhập tay (theo chuyến)"
    R.tien_di_lai = mileage_amount(total_km)

    # xăng: chỉnh hóa đơn CUỐI cho tổng = tiền đi lại theo KM (giống chế độ đơn giản)
    goc_last = gas_lines[-1]["thanh_toan"]
    tong_khac = sum(g["thanh_toan"] for g in gas_lines[:-1])
    can_co = int(round(R.tien_di_lai - tong_khac))
    if can_co < 0:
        err(f"Các hóa đơn xăng TRƯỚC hóa đơn cuối ({tong_khac:,.0f}) đã vượt tiền đi lại theo KM "
            f"({R.tien_di_lai:,.0f}). Bớt hóa đơn hoặc kiểm tra KM.")
        return R
    if can_co < goc_last:
        gas_lines[-1]["thanh_toan"] = can_co
        gas_lines[-1]["chua_vat"] = can_co - gas_lines[-1]["vat"]
        log(f"Hóa đơn xăng cuối HĐ {gas_lines[-1]['so_hd']}: {goc_last:,.0f} → {can_co:,.0f} (chỉnh cho khớp KM).")
    elif can_co > goc_last:
        R.thieu = can_co - goc_last
        warn(f"THIẾU HÓA ĐƠN XĂNG: còn thiếu {R.thieu:,.0f}đ so với tiền đi lại theo KM. "
             f"Thêm hóa đơn xăng rồi tạo lại.")
    gas_lines[-1]["is_last"] = True

    R.tien_ks = sum(h["thanh_toan"] for h in hotel_lines)
    pd_days = sum(t.so_dem for t in trips)
    R.tien_pd = pd_days * L.pd_rate
    R.tong_cp = R.tien_di_lai + R.tien_ks + R.tien_pd
    R.thu_chi_them = R.tong_cp - cfg.tam_ung
    R.tong_hd = sum(g["thanh_toan"] for g in gas_lines)
    R.khop = abs(R.tong_hd - R.tien_di_lai) < 1

    pd_line = {"so_hd": "AP4", "loai_hd": None, "mota": "Per-diem",
               "chua_vat": R.tien_pd, "vat": 0, "thanh_toan": R.tien_pd,
               "web": "", "ma_tra_cuu": "", "mst_ban": "", "trip_idx": None, "is_last": False}
    all_lines = gas_lines + hotel_lines + [pd_line]
    if len(all_lines) > L.dntt_inv_slots:
        err(f"{len(all_lines)} dòng hóa đơn nhưng mẫu DNTT chỉ {L.dntt_inv_slots} dòng. "
            f"Bớt hóa đơn hoặc dùng mẫu có nhiều dòng hơn.")
        return R
    R.invoices = all_lines

    # hotel theo chuyến (cho bảng claim)
    per_trip_hotel = {}
    for h in hotel_lines:
        per_trip_hotel[h["trip_idx"]] = per_trip_hotel.get(h["trip_idx"], 0) + h["thanh_toan"]

    if not khu_vuc:
        khu_vuc = trips[0].dia_diem
    try:
        R.xlsx_bytes = _fill_multi(wb, L, cfg, trips, all_lines, per_trip_hotel, khu_vuc)
    except Exception as ex:
        import traceback
        err(f"Lỗi khi điền Excel: {ex}\n{traceback.format_exc()[-500:]}")
        return R

    nm = _strip_diacritics(L.requestor) or "CS"
    R.out_filename = f"CS {nm} - {MON[cfg.month]} {cfg.year}.xlsx"
    R.ok = True
    return R


def _fill_multi(wb, L, cfg, trips, lines, per_trip_hotel, khu_vuc) -> bytes:
    dn = wb["DNTT"]
    dn["G3"] = f"Thời gian/ Time: {cfg.thang_label}"
    dn[f"C{L.dntt_desc_row0}"] = f"Công tác phí {khu_vuc} tháng {cfg.thang_label}"

    # xóa vùng hóa đơn cũ
    for r in range(L.dntt_inv_row0, L.dntt_sum_row):
        for col in ("A", "B", "F", "G", "H", "I", "J"):
            dn[f"{col}{r}"].value = None
            dn[f"{col}{r}"].hyperlink = None
    proj = str(dn[f"B{L.dntt_desc_row0}"].value or "").strip()
    for i, ln in enumerate(lines):
        r = L.dntt_inv_row0 + i
        dn[f"A{r}"] = i + 1
        if i == 0 and proj:
            dn[f"D{r}"] = proj
        dn[f"B{r}"] = ln["mota"]
        dn[f"F{r}"] = cfg.thang_label
        dn[f"G{r}"] = ln["chua_vat"]
        dn[f"H{r}"] = ln["vat"]
        dn[f"I{r}"] = ln["thanh_toan"]
        dn[f"J{r}"] = ln["so_hd"]

    # --- Bảng kê chi tiết ---
    bk = wb[L.bk_sheet]
    bk[L.bk_h3] = f"Thời gian/ Time: {cfg.month:02d}/{cfg.year}"
    if cfg.muc_dich:
        bk[L.bk_muc_dich_cell] = cfg.muc_dich

    for i in range(L.bk_trip_slots):
        r = L.bk_trip_row0 + i
        if i < len(trips):
            t = trips[i]
            bk[f"C{r}"] = t.dia_diem
            bk[f"D{r}"] = t.range_str
            bk[f"E{r}"] = t.so_dem
        else:
            bk[f"C{r}"] = None
            bk[f"D{r}"] = None
            bk[f"E{r}"] = 0

    for i in range(L.bk_claim_slots):
        r = L.bk_claim_row0 + i
        if i < len(trips):
            bk[f"D{r}"] = round(trips[i].km)
            hv = per_trip_hotel.get(i)
            bk[f"F{r}"] = hv if hv else None
        else:
            bk[f"D{r}"] = None
            bk[f"F{r}"] = None

    bk[L.bk_tam_ung_cell] = cfg.tam_ung if cfg.tam_ung else None

    # --- Tra cứu HĐ ---
    tc = wb[L.tc_sheet]
    for r in range(L.tc_row0, L.tc_row0 + 40):
        for col in ("B", "C", "D", "E"):
            tc[f"{col}{r}"].value = None
            tc[f"{col}{r}"].hyperlink = None
    r = L.tc_row0
    for ln in lines:
        if not ln["loai_hd"]:
            continue
        tc[f"B{r}"] = ln["loai_hd"]
        tc[f"C{r}"] = ln["web"]
        tc[f"D{r}"] = ln["ma_tra_cuu"]
        tc[f"E{r}"] = ln["mst_ban"] or ln["so_hd"]
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
